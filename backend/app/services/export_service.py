"""XLSX export builder — gaps / kb-drafts / marketing-brief review workbooks.

Mirrors the OLD Next.js exports (web/lib/gap-sheet.ts + web/app/api/.../export):
builds openpyxl workbooks from live DB data (gaps joined to run+ticket, personas)
plus the offline theme clusters artifact. Returns raw .xlsx bytes for streaming.
"""

from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.repositories.gap_repo import GapRepository, PersonaRepository

REVIEW_STATUS = ["Needs review", "Approved", "Edit requested", "Rejected"]
PUBLISH_READY = ["No", "Yes"]


def _outputs_dir() -> Path:
    return Path(__file__).resolve().parents[2] / "outputs"


def _today() -> str:
    return date.today().isoformat()


def _theme_clusters() -> list[dict]:
    path = _outputs_dir() / "theme_clusters.json"
    if not path.exists():
        return []
    try:
        return json.loads(path.read_text()) or []
    except (ValueError, OSError):
        return []


def _write_sheet(
    ws: Worksheet,
    rows: list[list[object]],
    widths: list[int] | None = None,
    validations: list[tuple[str, list[str]]] | None = None,
) -> None:
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    for sqref, options in validations or []:
        dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(sqref)


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ExportService:
    """Reads live data + offline artifacts, emits review-ready .xlsx bytes."""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def _open_gaps(self) -> list[dict]:
        rows = await GapRepository(self.session).list_enriched("open", limit=1000)
        gaps: list[dict] = []
        for gap, run, ticket in rows:
            first_seen = (ticket.date_received if ticket else None) or gap.created_at.date().isoformat()
            gaps.append(
                {
                    "ticket_id": ticket.ticket_id if ticket else str(gap.ticket_id),
                    "date_first_seen": first_seen,
                    "paraphrase": gap.paraphrase or "",
                    "theme": gap.theme or "",
                    "buyer_persona": gap.buyer_persona or "",
                    "status": gap.status,
                    "kb_confidence": run.kb_confidence if run else None,
                    "kb_entry_draft": gap.kb_entry_draft or "",
                }
            )
        return gaps

    # -- gaps.xlsx --------------------------------------------------------
    async def gaps_workbook(self) -> bytes:
        gaps = await self._open_gaps()
        header = [
            "Ticket ID",
            "Question",
            "Theme",
            "Persona",
            "Status",
            "KB Entry Draft",
        ]
        rows: list[list[object]] = [header]
        for gap in gaps:
            rows.append(
                [
                    gap["ticket_id"],
                    gap["paraphrase"],
                    gap["theme"],
                    gap["buyer_persona"],
                    gap["status"],
                    gap["kb_entry_draft"],
                ]
            )
        wb = Workbook()
        _write_sheet(wb.active, rows, widths=[18, 60, 22, 18, 14, 82])
        wb.active.title = "Knowledge Gaps"
        return _save(wb)

    # -- kb-drafts.xlsx ---------------------------------------------------
    async def kb_drafts_workbook(self) -> bytes:
        gaps = [g for g in await self._open_gaps() if g["kb_entry_draft"]]
        header = [
            "Approval Status",
            "Publish Ready",
            "Source Ticket ID",
            "Question",
            "Theme",
            "Persona",
            "KB Entry Draft",
            "Approver Notes",
        ]
        rows: list[list[object]] = [header]
        for gap in gaps:
            rows.append(
                [
                    "Needs review",
                    "No",
                    gap["ticket_id"],
                    gap["paraphrase"],
                    gap["theme"],
                    gap["buyer_persona"],
                    gap["kb_entry_draft"],
                    "",
                ]
            )
        wb = Workbook()
        _write_sheet(
            wb.active,
            rows,
            widths=[18, 14, 18, 58, 22, 18, 82, 44],
            validations=[("A2:A1000", REVIEW_STATUS), ("B2:B1000", PUBLISH_READY)],
        )
        wb.active.title = "KB Drafts"
        return _save(wb)

    # -- marketing-brief.xlsx --------------------------------------------
    async def marketing_brief_workbook(self) -> bytes:
        gaps = await self._open_gaps()
        personas = await PersonaRepository(self.session).list(limit=50)
        themes = _theme_clusters()
        wb = Workbook()
        wb.remove(wb.active)
        self._brief_overview(wb, gaps, personas)
        self._brief_gaps(wb, gaps)
        self._brief_kb_drafts(wb, gaps)
        self._brief_themes(wb, themes)
        self._brief_personas(wb, personas)
        return _save(wb)

    def _brief_overview(self, wb: Workbook, gaps: list[dict], personas) -> None:
        persona_counts: dict[str, int] = {}
        for gap in gaps:
            key = gap["buyer_persona"] or "?"
            persona_counts[key] = persona_counts.get(key, 0) + 1
        rows: list[list[object]] = [
            ["Boldr Marketing Brief Export", ""],
            ["Generated", _today()],
            [
                "Purpose",
                "Human review of AI-generated segmentation, campaign angles, and KB feedback outputs.",
            ],
            ["Knowledge Gaps", len(gaps)],
            ["Personas", len(personas)],
            ["Review Status Values", " | ".join(REVIEW_STATUS)],
            ["", ""],
            ["Persona", "Gap Count"],
        ]
        rows.extend([persona, count] for persona, count in persona_counts.items())
        _write_sheet(wb.create_sheet("Overview"), rows, widths=[28, 90])

    def _brief_gaps(self, wb: Workbook, gaps: list[dict]) -> None:
        header = [
            "Review Status",
            "Ticket ID",
            "Persona",
            "Theme",
            "Question",
            "KB Entry Draft",
            "Human Notes",
        ]
        rows: list[list[object]] = [header]
        for gap in gaps:
            rows.append(
                [
                    "Needs review",
                    gap["ticket_id"],
                    gap["buyer_persona"],
                    gap["theme"],
                    gap["paraphrase"],
                    gap["kb_entry_draft"],
                    "",
                ]
            )
        _write_sheet(
            wb.create_sheet("Knowledge Gaps"),
            rows,
            widths=[16, 18, 18, 20, 58, 82, 44],
            validations=[("A2:A1000", REVIEW_STATUS)],
        )

    def _brief_kb_drafts(self, wb: Workbook, gaps: list[dict]) -> None:
        header = [
            "Approval Status",
            "Publish Ready",
            "Source Ticket ID",
            "Question",
            "Theme",
            "Persona",
            "KB Entry Draft",
            "Approver Notes",
        ]
        rows: list[list[object]] = [header]
        for gap in (g for g in gaps if g["kb_entry_draft"]):
            rows.append(
                [
                    "Needs review",
                    "No",
                    gap["ticket_id"],
                    gap["paraphrase"],
                    gap["theme"],
                    gap["buyer_persona"],
                    gap["kb_entry_draft"],
                    "",
                ]
            )
        _write_sheet(
            wb.create_sheet("KB Drafts"),
            rows,
            widths=[18, 14, 18, 58, 22, 18, 82, 44],
            validations=[("A2:A1000", REVIEW_STATUS), ("B2:B1000", PUBLISH_READY)],
        )

    def _brief_themes(self, wb: Workbook, themes: list[dict]) -> None:
        header = [
            "Theme",
            "Cluster Size",
            "Persona Mix",
            "Marketing Signal",
            "Suggested Action",
            "Sample Questions",
            "Human Priority",
            "Owner",
            "Notes",
        ]
        rows: list[list[object]] = [header]
        for theme in (t for t in themes if t.get("size", 0) > 1):
            persona_mix = ", ".join(
                f"{persona}: {count}"
                for persona, count in (theme.get("persona_breakdown") or {}).items()
            )
            rows.append(
                [
                    theme.get("theme_label", ""),
                    theme.get("size", 0),
                    persona_mix,
                    theme.get("marketing_signal", ""),
                    theme.get("suggested_action", ""),
                    " | ".join(theme.get("sample_questions") or []),
                    "",
                    "",
                    "",
                ]
            )
        _write_sheet(
            wb.create_sheet("Theme Actions"),
            rows,
            widths=[34, 12, 30, 70, 70, 58, 18, 20, 44],
        )

    def _brief_personas(self, wb: Workbook, personas) -> None:
        header = [
            "Persona ID",
            "Persona Name",
            "Priority",
            "Trigger Keywords",
            "Recommended Messaging",
            "Marketing Opportunity",
        ]
        rows: list[list[object]] = [header]
        for persona in personas:
            rows.append(
                [
                    persona.persona_id,
                    persona.name,
                    persona.priority,
                    ", ".join(persona.trigger_keywords or []),
                    persona.recommended_messaging,
                    persona.marketing_opportunity,
                ]
            )
        _write_sheet(
            wb.create_sheet("Persona Reference"),
            rows,
            widths=[20, 28, 14, 68, 70, 70],
        )
